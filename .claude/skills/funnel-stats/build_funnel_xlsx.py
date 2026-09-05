#!/usr/bin/env python3
"""Сборка funnel-stats.xlsx с нуля: таблица дневной воронки + диаграммы + анализ.

Вход — результат SQL-запроса к daily_stats со stdin, по строке на день,
поля разделены '|' в порядке:
    date|group_replies|public_subs_abs|post_views_abs|post_views_delta|private_requests_delta
Пустое числовое поле трактуется как 0 (дельты первого дня = NULL в БД).
post_views_abs из запроса в таблицу не выводится (нужен только прирост).

Лист «Данные»:
  A–E — сырьё/притоки из daily_stats;
  F–H — конверсия просмотр/ответ и её скользящее среднее (СС-3);
  «Притоки заявок» — положительная часть прироста заявок (приём в канал, т.е.
  отрицательные дельты, не считается спросом и обнуляется).
Первый день — базовый (дельта = 0 как артефакт отсутствия предыдущего снапшота):
его конверсия не считается (пусто) и он исключён из статистики.

Под таблицей — блок «Анализ за период»: корреляция Пирсона и линейная регрессия
«прирост просмотров ~ ответы» одним числом за весь диапазон, период, число дней,
уравнение и средняя конверсия.

Диаграммы (правее данных, сверху вниз): «Конверсия просмотр/ответ + тренд» →
столбчатые: ответы, подписчики, прирост просмотров, притоки заявок. У каждой
подписаны обе оси (название + числовые деления); чисел на столбцах/точках нет.

Лист «Описание»: что за каждый производный параметр, зачем и как вычисляется.

Использование:
    ssh ... 'sqlite3 ... "<query>"' | python build_funnel_xlsx.py <out.xlsx>
"""
import sys

import openpyxl
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.chart.trendline import Trendline
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

# Порядок столбцов листа «Данные».
HEADERS = [
    "Дата",
    "Ответы в группах (Σ)",
    "Подписчики (Σ abs)",
    "Прирост просмотров (Σ)",
    "Прирост заявок в приват (Σ)",
    "Притоки заявок (Σ)",
    "Конверсия просмотр/ответ",
    "Конверсия СС-3",
]

COL_REPLIES = 2      # B
COL_SUBS = 3         # C
COL_VIEWS_DELTA = 4  # D
COL_PRIV = 5         # E — прирост заявок (нетто, со знаком)
COL_INFLOW = 6       # F — притоки заявок (только положительная часть)
COL_CONV = 7         # G — дневная конверсия
COL_CONV_MA = 8      # H — скользящее среднее конверсии
MA_WINDOW = 3        # окно скользящего среднего конверсии (дни)

# Столбчатые диаграммы: (колонка, заголовок, подпись оси Y).
BAR_CHARTS = [
    (COL_REPLIES, "Ответы в группах по дням", "Ответы, шт"),
    (COL_SUBS, "Подписчики (абсолют) по дням", "Подписчики, чел"),
    (COL_VIEWS_DELTA, "Прирост просмотров по дням", "Просмотры/день, шт"),
    (COL_INFLOW, "Притоки заявок в приват по дням", "Заявки/день, шт"),
]

# Габариты диаграммы (см) и вертикальный шаг якорей (строки) — без наложения.
CHART_W, CHART_H = 14, 7
CHART_ANCHOR_COL = get_column_letter(len(HEADERS) + 2)  # правее данных + столбец-зазор
CHART_ROW_STEP = 15
BAR_COLOR = "4472C4"  # единый цвет столбцов на всех диаграммах
AXIS_X_TITLE = "Дата"

# Описания производных параметров: (параметр, что это, зачем, как вычисляется).
PARAM_DESCRIPTIONS = [
    (
        "Притоки заявок в приват",
        "Сколько новых заявок в приватный канал поступило за день (только положительная часть изменения).",
        "Очищает спрос от приёма в канал. Связан с приростом просмотров (r≈0.73): больше просмотров — больше заявок.",
        "Положительная часть «Прирост заявок в приват (Σ)»: само значение, если оно > 0, иначе 0. В дни приёма (нетто < 0) истинный приток замаскирован и показан 0.",
    ),
    (
        "Конверсия просмотр/ответ",
        "Сколько новых просмотров публичного канала приходится на один ответ агента в группе за день.",
        "Главный показатель эффективности: растёт охват с каждого ответа или воронка насыщается. Падение при росте ответов = новые ответы дают меньше отдачи.",
        "Прирост просмотров (Σ) ÷ Ответы в группах (Σ) за тот же день. Базовый день и дни без ответов не считаются.",
    ),
    (
        "Конверсия СС-3",
        "Скользящее среднее дневной конверсии за 3 последних дня.",
        "Убирает дневной шум (выходные, всплески) и показывает устойчивый тренд конверсии, а не случайные скачки.",
        "Среднее значение «Конверсия просмотр/ответ» за текущий и два предыдущих дня; пока трёх значений нет — пусто.",
    ),
    (
        "Корреляция Пирсона (ответы↔прирост просмотров)",
        "Сила линейной связи между числом ответов и приростом просмотров, от −1 до +1.",
        "Проверяет, действительно ли просмотры тянутся за ответами. Близко к +1 — связь сильная; около 0 — просмотры от ответов почти не зависят.",
        "Ковариация ответов и прироста просмотров ÷ произведение их стандартных отклонений (по дням без базового).",
    ),
    (
        "R² (доля объяснённой дисперсии)",
        "Доля колебаний прироста просмотров, которую объясняет число ответов, от 0 до 1.",
        "Показывает, насколько модель «просмотры зависят от ответов» полна. 0.69 = 69% разброса объясняется ответами, 31% — прочие факторы (органика, дни недели).",
        "Квадрат корреляции Пирсона.",
    ),
    (
        "Регрессия: базовый уровень a",
        "Прогноз прироста просмотров в день, когда ответов ноль.",
        "Оценивает «органический фон» — просмотры, приходящие без ответов (поиск, старые посты, репосты).",
        "Свободный член прямой y = a + b·x, подобранной методом наименьших квадратов (x — ответы, y — прирост просмотров).",
    ),
    (
        "Регрессия: отдача на ответ b",
        "На сколько в среднем растёт прирост просмотров при +1 ответе.",
        "Реальная «цена» одного ответа в просмотрах. Снижение b со временем — прямое доказательство насыщения воронки.",
        "Наклон прямой y = a + b·x (метод наименьших квадратов).",
    ),
    (
        "Средняя конверсия просмотр/ответ",
        "Среднее значение дневной конверсии за весь период.",
        "Ориентир «нормы» отдачи с ответа для сравнения отдельных дней с типичным уровнем.",
        "Среднее арифметическое всех непустых значений «Конверсия просмотр/ответ».",
    ),
]


def parse_stdin(text):
    """pipe-separated строки -> список кортежей (date, replies, subs, views_abs,
    views_delta, priv_delta) в порядке исходного SQL."""
    rows = []
    text = text.lstrip("﻿")  # срезать BOM, если файл с ним
    for line in text.splitlines():
        line = line.strip()
        if not line:
            continue
        parts = line.split("|")
        if len(parts) != 6:
            raise ValueError(f"ожидалось 6 полей, получено {len(parts)}: {line!r}")
        date = parts[0]
        nums = [int(p) if p.strip() else 0 for p in parts[1:]]
        rows.append((date, *nums))
    if not rows:
        raise ValueError("нет данных на stdin")
    return rows


def conversions(data):
    """Дневная конверсия просмотр/ответ. Первый день (базовый) и дни без ответов
    -> None: делить не на что / делёж даёт артефакт."""
    conv = []
    for i, (_date, replies, _subs, _va, views_delta, _priv) in enumerate(data):
        if i == 0 or replies == 0:
            conv.append(None)
        else:
            conv.append(round(views_delta / replies, 2))
    return conv


def moving_average(values, window):
    """Скользящее среднее по непустым значениям; пока окно не заполнено — None."""
    out = []
    for i in range(len(values)):
        chunk = values[max(0, i - window + 1): i + 1]
        nums = [v for v in chunk if v is not None]
        out.append(round(sum(nums) / len(nums), 2) if len(nums) == window else None)
    return out


def pearson(xs, ys):
    """Корреляция Пирсона; None при <2 точках или нулевой дисперсии."""
    n = len(xs)
    if n < 2:
        return None
    mx, my = sum(xs) / n, sum(ys) / n
    sx = sum((x - mx) ** 2 for x in xs) ** 0.5
    sy = sum((y - my) ** 2 for y in ys) ** 0.5
    if sx == 0 or sy == 0:
        return None
    cov = sum((x - mx) * (y - my) for x, y in zip(xs, ys))
    return cov / (sx * sy)


def linreg(xs, ys):
    """МНК: y = a + b*x. (None, None) при <2 точках или нулевой дисперсии x."""
    n = len(xs)
    if n < 2:
        return None, None
    mx, my = sum(xs) / n, sum(ys) / n
    denom = sum((x - mx) ** 2 for x in xs)
    if denom == 0:
        return None, None
    b = sum((x - mx) * (y - my) for x, y in zip(xs, ys)) / denom
    return my - b * mx, b


def _style_axes(ch, y_title):
    """Подписать обе оси (название) и показать числовые деления на них."""
    ch.x_axis.title = AXIS_X_TITLE
    ch.y_axis.title = y_title
    ch.x_axis.delete = False  # иначе Excel скрывает ось и её числа
    ch.y_axis.delete = False


def build(data, out_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Данные"

    conv = conversions(data)
    conv_ma = moving_average(conv, MA_WINDOW)

    ws.append(HEADERS)
    for i, (date, replies, subs, _views_abs, views_delta, priv) in enumerate(data):
        inflow = priv if priv > 0 else 0  # притоки: только положительная часть
        ws.append([date, replies, subs, views_delta, priv, inflow, conv[i], conv_ma[i]])

    # Оформление шапки и колонок.
    for c in range(1, len(HEADERS) + 1):
        cell = ws.cell(row=1, column=c)
        cell.font = Font(bold=True)
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[1].height = 42
    ws.column_dimensions["A"].width = 13
    for c in range(2, len(HEADERS) + 1):
        ws.column_dimensions[get_column_letter(c)].width = 16
    ws.freeze_panes = "A2"

    last = len(data) + 1  # последняя строка данных (строка 1 — шапка)
    cats = Reference(ws, min_col=1, min_row=2, max_row=last)

    anchor_row = 1

    # 1) Конверсия — первой: дневная + СС-3 + линейная линия тренда.
    lc = LineChart()
    lc.title = "Конверсия просмотр/ответ + тренд"
    lc.width, lc.height = CHART_W, CHART_H
    lc.add_data(Reference(ws, min_col=COL_CONV, min_row=1, max_row=last), titles_from_data=True)
    lc.add_data(Reference(ws, min_col=COL_CONV_MA, min_row=1, max_row=last), titles_from_data=True)
    lc.set_categories(cats)
    _style_axes(lc, "Просмотров на ответ")
    lc.series[0].trendline = Trendline(trendlineType="linear", dispEq=True, dispRSqr=True)
    ws.add_chart(lc, f"{CHART_ANCHOR_COL}{anchor_row}")
    anchor_row += CHART_ROW_STEP

    # 2) Столбчатые диаграммы — ниже.
    for col, title, y_title in BAR_CHARTS:
        ch = BarChart()
        ch.type = "col"
        ch.title = title
        ch.width, ch.height = CHART_W, CHART_H
        ch.legend = None
        ch.varyColors = False  # иначе каждый столбец красится в свой цвет
        ch.add_data(Reference(ws, min_col=col, min_row=1, max_row=last), titles_from_data=True)
        ch.set_categories(cats)
        ch.series[0].graphicalProperties = GraphicalProperties(solidFill=BAR_COLOR)
        _style_axes(ch, y_title)
        ws.add_chart(ch, f"{CHART_ANCHOR_COL}{anchor_row}")
        anchor_row += CHART_ROW_STEP

    _write_analysis_block(ws, data, start_row=last + 2)
    _build_descriptions(wb)

    wb.save(out_path)
    return last - 1  # число строк данных


def _write_analysis_block(ws, data, start_row):
    """Блок «Анализ за период» под таблицей (колонки A — подпись, C — значение).
    Базовый (первый) день исключён — его дельта = 0 как артефакт, не реальный ноль."""
    xs = [r[1] for r in data[1:]]  # ответы
    ys = [r[4] for r in data[1:]]  # прирост просмотров

    title = ws.cell(row=start_row, column=1, value="Анализ за период (все дни, кроме базового)")
    title.font = Font(bold=True)

    if len(xs) >= 2:
        r = pearson(xs, ys)
        a, b = linreg(xs, ys)
        convs = [v for v in conversions(data) if v is not None]
        rows = [
            ("Период", f"{data[0][0]}…{data[-1][0]}"),
            ("Дней в анализе (без базового)", len(xs)),
            ("Корреляция Пирсона", round(r, 3) if r is not None else "—"),
            ("R²", round(r * r, 3) if r is not None else "—"),
            ("Регрессия: фон a", round(a, 1) if a is not None else "—"),
            ("Регрессия: отдача b", round(b, 3) if b is not None else "—"),
            ("Уравнение", f"прирост просмотров = {a:.1f} + {b:.2f}·ответы" if a is not None else "—"),
            ("Средняя конверсия", round(sum(convs) / len(convs), 2) if convs else "—"),
        ]
    else:
        rows = [("Недостаточно дней для анализа", len(xs))]

    for i, (label, value) in enumerate(rows):
        row = start_row + 1 + i
        ws.cell(row=row, column=1, value=label)  # B пуст → подпись перетекает вправо
        ws.cell(row=row, column=3, value=value)


def _build_descriptions(wb):
    """Лист «Описание»: по каждому производному параметру — что это, зачем, как считается."""
    ws = wb.create_sheet("Описание")
    headers = ["Параметр", "Что это", "Зачем нужен (что определяет)", "Как вычисляется"]
    ws.append(headers)
    for row in PARAM_DESCRIPTIONS:
        ws.append(list(row))
    for c in range(1, len(headers) + 1):
        ws.cell(row=1, column=c).font = Font(bold=True)
        ws.cell(row=1, column=c).alignment = Alignment(wrap_text=True, vertical="center")
    ws.column_dimensions["A"].width = 30
    for col in ("B", "C", "D"):
        ws.column_dimensions[col].width = 52
    for r in range(2, len(PARAM_DESCRIPTIONS) + 2):
        for c in range(2, 5):
            ws.cell(row=r, column=c).alignment = Alignment(wrap_text=True, vertical="top")


def main():
    if len(sys.argv) != 2:
        sys.exit("usage: build_funnel_xlsx.py <out.xlsx>")
    data = parse_stdin(sys.stdin.read())
    n = build(data, sys.argv[1])
    print(f"saved {sys.argv[1]}: {n} строк ({data[0][0]}…{data[-1][0]}), "
          f"{len(BAR_CHARTS) + 1} диаграмм + блок «Анализ» под таблицей + лист «Описание»")


if __name__ == "__main__":
    main()
